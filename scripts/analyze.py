#!/usr/bin/env python3
"""Mobile proxy benchmark: merge the CSVs of the three servers into one workbook.

  python3 -I analyze.py --raw RAWDIR --out results.xlsx [--ipq IPQDIR]

RAWDIR contains one folder per server (us, de, sg), each with the folder data/ from that server.
IPQDIR (optional) contains ipsample.csv and ipcache.json from the IP-quality run (ipquality.py).
Needs pandas and openpyxl.
"""
import argparse, csv, glob, json, math, os, statistics, sys
from collections import Counter, defaultdict

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import sticky  # noqa: E402  (IP lifetime rule)
import ipquality  # noqa: E402  (parse of the Scamalytics / IPinfo answers)
import fingerprint  # noqa: E402  (OS groups)

START = "2026-10-24 00:00:00"          # nothing before the official start counts (the pilot is in pilot/)
STICKY_MIN = 30
SERVERS = ["us", "de", "sg"]
SERVER_NAMES = {"us": "USA (Ashburn)", "de": "Germany (Falkenstein)", "sg": "Singapore"}
COUNTRIES = ["us", "de", "gb", "fr", "es", "br", "in", "pl", "ua", "tr", "vn", "id"]

LABEL = {"floppydata": "Floppydata", "decodo": "Decodo", "oxylabs": "Oxylabs", "evomi": "Evomi", "soax": "SOAX",
         "dataimpulse": "DataImpulse", "proxyempire": "ProxyEmpire", "nodemaven": "NodeMaven", "iproyal": "IPRoyal",
         "byteful": "Byteful", "resiprox": "ResiProx"}
ORDER = ["floppydata", "byteful", "dataimpulse", "decodo", "evomi", "iproyal", "nodemaven", "oxylabs", "proxyempire",
         "resiprox", "soax"]
# regular price of the smallest package, $/GB, without promo codes or VAT (checked 23-26 Oct 2026), and the
# smallest self-serve purchase. Re-check both right before publication.
PRICE = {"floppydata": (1.00, "19 GB for $19"),
         "dataimpulse": (2.00, "2.5 GB for $5, one-time (new users)"),
         "decodo": (3.75, "2 GB for $7.50, subscription"),
         "evomi": (3.98, "4 GB for $15.92, one-time"),
         "byteful": (4.25, "2 GB for $8.50, subscription"),
         "soax": (4.60, "$25 top-up (25 credits, expire in 60 days)"),  # effective: $5 / $4 / $2 by country tier
         "nodemaven": (5.00, "2 GB for $10, one-time"),
         "resiprox": (6.00, "1 GB for $6, subscription"),
         "oxylabs": (7.50, "4 GB for $30, subscription only"),
         "iproyal": (7.65, "2 GB for $15.30, subscription"),
         "proxyempire": (9.00, "1 GB for $9, one-time")}


def wilson(k, n, z=1.96):
    if n == 0:
        return (float("nan"), float("nan"))
    p = k / n
    d = 1 + z * z / n
    c = p + z * z / (2 * n)
    m = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n))
    return ((c - m) / d, (c + m) / d)


def fmt_rate(k, n):
    if n == 0:
        return "-"
    lo, hi = wilson(k, n)
    return "%.1f%% (%.1f-%.1f)" % (100.0 * k / n, max(0.0, 100 * lo), min(100.0, 100 * hi))


def norm_ts(s):
    return s.replace("T", " ").replace("Z", "")


def read_all(raw, prefix):
    frames = []
    for srv in SERVERS:
        for fn in sorted(glob.glob(os.path.join(raw, srv, "data", prefix + "_*.csv"))):
            df = pd.read_csv(fn, dtype=str, keep_default_na=False)
            if len(df):
                df["server"] = srv
                frames.append(df)
    if not frames:
        return pd.DataFrame()
    df = pd.concat(frames, ignore_index=True)
    df["ts"] = df["ts_utc"].map(norm_ts)
    return df[df["ts"] >= START].copy()


def num(df, cols):
    for c in cols:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    return df


# ------------------------------------------------------------------ analyses
def load_infra(raw):
    df = read_all(raw, "infra")
    if df.empty:
        return df
    df = num(df, ["ok", "sec", "bytes"])
    df["geo_ok"] = (df["geo_cf"].str.lower() == df["country_req"].str.lower()) & (df["ok"] == 1)
    # the ipv column, not the address: the public copy has hashed addresses
    df["v6"] = (df["ipv"].astype(str) == "6") & (df["ok"] == 1)
    return df


def infra_stats(g):
    n = len(g)
    ok = g[g["ok"] == 1]
    s = sorted(ok["sec"].tolist())
    uniq = ok["ip"].nunique()
    return dict(n=n, ok=len(ok), success=len(ok) / n if n else float("nan"),
                lo=wilson(len(ok), n)[0], hi=wilson(len(ok), n)[1],
                median=statistics.median(s) if s else float("nan"),
                p90=s[int(0.9 * (len(s) - 1))] if s else float("nan"),
                unique=uniq, dup=(1 - uniq / len(ok)) if len(ok) else float("nan"),
                geo=(ok["geo_ok"].sum() / len(ok)) if len(ok) else float("nan"),
                geo_k=int(ok["geo_ok"].sum()), v6=(ok["v6"].sum() / len(ok)) if len(ok) else float("nan"))


def load_targets(raw):
    df = read_all(raw, "targets")
    if df.empty:
        return df
    return num(df, ["ok", "sec", "bytes"])


def target_stats(g):
    g = g[g["result"] != "provider_restricted"]
    n = len(g)
    k = int((g["ok"] == 1).sum())
    cap = int((g["result"] == "captcha").sum())
    blk = int((g["result"].str.startswith("blocked_") | g["result"].isin(["rate_limited", "login_wall"])).sum())
    okg = g[g["ok"] == 1]
    return dict(n=n, k=k, cap=cap, blk=blk, oth=n - k - cap - blk,
                med=okg["sec"].median() if len(okg) else float("nan"),
                p90=okg["sec"].quantile(0.9) if len(okg) else float("nan"),
                kb=g["bytes"].mean() / 1024 if n else float("nan"))


def sticky_sessions(raw):
    df = read_all(raw, "sticky")
    if df.empty:
        return []
    out = []
    for (run, srv, prov, ses), g in df.groupby(["run_id", "server", "provider", "session"]):
        rs = g.to_dict("records")
        life, full = sticky.lifetime(rs, STICKY_MIN)
        out.append(dict(run=run, server=srv, provider=prov, started=life is not None, life=life or 0.0, full=full,
                        n=len(rs), ok=sum(1 for r in rs if r["ok"] == "1")))
    return out


def sticky_stats(ss):
    n = len(ss)
    k10 = sum(s["life"] >= 10 for s in ss)
    kf = sum(s["full"] for s in ss)
    okreq, allreq = sum(s["ok"] for s in ss), sum(s["n"] for s in ss)
    return dict(n=n, k10=k10, kf=kf, med=statistics.median(s["life"] for s in ss) if ss else float("nan"),
                r10=k10 / n if n else float("nan"), rf=kf / n if n else float("nan"),
                lo=wilson(kf, n)[0], hi=wilson(kf, n)[1],
                err=((allreq - okreq) / allreq) if allreq else float("nan"), okreq=okreq)


def load_fp(raw):
    df = read_all(raw, "fingerprint")
    if df.empty:
        return df
    df = num(df, ["ok", "sec"])
    df["os_group"] = [fingerprint.os_group(r) for r in df.to_dict("records")]
    df["eth"] = df["link"].str.lower().str.startswith("ethernet")
    df["link_known"] = (df["link"] != "") & (df["link"] != "???")
    return df


def fp_stats(g):
    ok = g[g["ok"] == 1]
    n = len(ok)
    c = Counter(ok["os_group"])
    lk = ok[ok["link_known"]]
    return dict(n=n, req=len(g), linux=c["linux"], windows=c["windows"], mac=c["mac os"], other=c["other"] + c["freebsd"],
                unknown=c["unknown"], eth=int(lk["eth"].sum()), link_n=len(lk),
                tunnel=int(lk["link"].str.contains("tunnel").sum()))


def ipq_stats(ipq):
    sample = os.path.join(ipq, "ipsample.csv")
    cache_p = os.path.join(ipq, "ipcache.json")
    if not (os.path.exists(sample) and os.path.exists(cache_p)):
        return {}
    cache = json.load(open(cache_p))
    by = defaultdict(list)
    for r in csv.DictReader(open(sample, newline="")):
        if r["ip"] in cache:
            by[r["provider"]].append(ipquality.parse(cache[r["ip"]], r["country_req"]))
    out = {}
    for p, ds in by.items():
        v = [d for d in ds if d["ok"]]
        if not v:
            continue
        gi = [d for d in ds if d["ipinfo_ok"]]
        out[p] = dict(n=len(v), avg=statistics.mean(d["score"] for d in v), med=statistics.median(d["score"] for d in v),
                      high=sum(d["risk"] in ("high", "very high") for d in v), proxy=sum(d["proxy"] for d in v),
                      dc=sum(d["dc"] for d in v), vpn=sum(d["vpn"] for d in v),
                      host=sum(d["host_asn"] for d in ds), all=len(ds), geo=sum(d["geo"] for d in gi), geo_n=len(gi))
    return out


def ip2l_mobile(u):
    return "MOB" in (u or "").split("/")


def mm_cellular(ans):
    t = ans.get("traits") or {}
    return t.get("user_type") == "cellular" or t.get("connection_type") == "Cellular"


def mobile_stats(ipq):
    """Mobile IP Share by IP2Location (usage type contains MOB) and by MaxMind (user type cellular).

    MaxMind checked the first 1,000 IPs per provider of the IP2Location sample, so 'both' and 'MOB only' are
    counted on that common subset."""
    out = defaultdict(lambda: Counter())
    paths = {k: (os.path.join(ipq, k + "sample.csv"), os.path.join(ipq, k + "cache.json")) for k in ("ip2l", "mm")}
    if not all(os.path.exists(x) for pair in paths.values() for x in pair):
        return {}
    c2, cm = json.load(open(paths["ip2l"][1])), json.load(open(paths["mm"][1]))
    for r in csv.DictReader(open(paths["ip2l"][0], newline="")):
        a = (c2.get(r["ip"]) or {}).get("ip2l")
        if a:
            out[r["provider"]]["ip2l_n"] += 1
            out[r["provider"]]["ip2l_k"] += ip2l_mobile(a.get("usage_type"))
    for r in csv.DictReader(open(paths["mm"][0], newline="")):
        b = (cm.get(r["ip"]) or {}).get("mm")
        if not b:
            continue
        o = out[r["provider"]]
        o["mm_n"] += 1
        o["mm_k"] += mm_cellular(b)
        a = (c2.get(r["ip"]) or {}).get("ip2l")
        if a:
            u = a.get("usage_type")
            o["both_n"] += 1
            o["both_k"] += ip2l_mobile(u) and mm_cellular(b)
            o["mobonly"] += u == "MOB"
            o["ip2l_not_mm"] += ip2l_mobile(u) and not mm_cellular(b)
    return dict(out)


def shared_ips(infra, provs):
    """Exit IPs of each provider that some other provider also gave out during the test."""
    ok = infra[infra["ok"] == 1]
    ips = {p: set(ok[ok["provider"] == p]["ip"]) for p in provs}
    hours = {p: ok[ok["provider"] == p].groupby("ip")["ts"].agg(lambda s: set(x[:13] for x in s)).to_dict() for p in provs}
    any_other = {}
    for p in provs:
        others = set().union(*(ips[q] for q in provs if q != p))
        any_other[p] = (len(ips[p] & others), len(ips[p]))
    pairs = []
    for i, a in enumerate(provs):
        for b in provs[i + 1:]:
            sh = ips[a] & ips[b]
            if sh:
                same = sum(1 for ip in sh if hours[a][ip] & hours[b][ip])
                pairs.append((a, b, len(sh), len(sh) / len(ips[a]), len(sh) / len(ips[b]), same / len(sh)))
    pairs.sort(key=lambda x: -x[2])
    return any_other, pairs


# ------------------------------------------------------------------ workbook
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")


def sheet(wb, title, header, rows, widths=None, pct_cols=(), num_cols=None, freeze="B2", note=None):
    ws = wb.create_sheet(title[:31])
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(italic=True, color="555555")
        r0 = 3
    for j, h in enumerate(header, 1):
        c = ws.cell(r0, j, h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[r0].height = 45
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, float) and math.isnan(v):
                v = None
            c = ws.cell(i, j, v)
            if j in pct_cols and isinstance(v, (int, float)):
                c.number_format = "0.0%"
            elif num_cols and j in num_cols and isinstance(v, (int, float)):
                c.number_format = num_cols[j]
    for j in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(j, 16)
    ws.column_dimensions["A"].width = (widths or {}).get(1, 20)
    ws.freeze_panes = ws.cell(r0 + 1, 2) if freeze else None
    return ws


def per_server(fn, data, provs, key="provider"):
    return {(p, s): fn(data[(data[key] == p) & (data["server"] == s)]) for p in provs for s in SERVERS}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw", required=True)
    ap.add_argument("--out", default="results.xlsx")
    ap.add_argument("--ipq")
    ap.add_argument("--json", help="also write per-provider numbers for the website here")
    ap.add_argument("--counters", help="CSV provider,used_mb,read_at_utc,source,note: dashboard traffic counters for the Cost sheet")
    a = ap.parse_args()
    counters = {}
    if a.counters:
        with open(a.counters, newline="") as fh:
            counters = {r["provider"]: r for r in csv.DictReader(fh)}

    infra = load_infra(a.raw)
    targets = load_targets(a.raw)
    stk = sticky_sessions(a.raw)
    fp = load_fp(a.raw)
    ipq = ipq_stats(a.ipq) if a.ipq else {}
    mob = mobile_stats(a.ipq) if a.ipq else {}
    seen = set(infra["provider"]) if len(infra) else set()
    provs = [p for p in ORDER if p in seen]
    shared_any, shared_pairs = shared_ips(infra, provs) if len(infra) else ({}, [])

    P = {p: infra_stats(infra[infra["provider"] == p]) for p in provs}
    PS = per_server(infra_stats, infra, provs)
    ST = {p: sticky_stats([s for s in stk if s["provider"] == p]) for p in provs}
    STS = {(p, sv): sticky_stats([s for s in stk if s["provider"] == p and s["server"] == sv]) for p in provs for sv in SERVERS}
    TG = {(p, t): target_stats(targets[(targets["provider"] == p) & (targets["target"] == t)])
          for p in provs for t in ("google", "instagram")} if len(targets) else {}
    FP = {p: fp_stats(fp[fp["provider"] == p]) for p in provs} if len(fp) else {}
    weakest = {}
    for p in provs:
        sr = infra[infra["provider"] == p].groupby("country_req")["ok"].mean()
        if len(sr):
            weakest[p] = "%s %.1f%%" % (sr.idxmin().upper(), 100 * sr.min())

    wb = Workbook()
    wb.remove(wb.active)
    period = "%s .. %s UTC" % (infra["ts"].min()[:16], infra["ts"].max()[:16]) if len(infra) else "-"

    # ---------------- Summary
    header = ["Provider", "Success Rate", "Latency median / P90, s", "Unique IPs", "Duplicate Rate", "Geo Accuracy",
              "IPv6 share", "Weakest Country", "Google Success Rate", "Instagram Success Rate",
              "Sticky: median IP lifetime, min", "Sticky: lived 30 min", "Fingerprint: Windows share",
              "Fingerprint: Ethernet MTU 1500", "Fraud Score avg / high risk", "Proxy flags (Scamalytics)",
              "Mobile IP Share", "Cost per 1,000 successful, $", "Price, $/GB", "Minimum purchase"]
    rows = []
    for p in provs:
        s, st = P[p], ST[p]
        g = TG.get((p, "google"), {"k": 0, "n": 0})
        ig = TG.get((p, "instagram"), {"k": 0, "n": 0})
        f = FP.get(p)
        q = ipq.get(p)
        rows.append([LABEL[p], fmt_rate(s["ok"], s["n"]), "%.2f / %.2f" % (s["median"], s["p90"]),
                     "%d of %d" % (s["unique"], s["ok"]), s["dup"], fmt_rate(s["geo_k"], s["ok"]), s["v6"],
                     weakest.get(p, "-"), fmt_rate(g["k"], g["n"]), fmt_rate(ig["k"], ig["n"]),
                     st["med"], fmt_rate(st["kf"], st["n"]) + " (n=%d)" % st["n"] if st["n"] else "-",
                     fmt_rate(f["windows"], f["n"]) if f and f["n"] else "-",
                     fmt_rate(f["eth"], f["link_n"]) if f and f["link_n"] else "-",
                     "%.1f / %s" % (q["avg"], fmt_rate(q["high"], q["n"])) if q else "pending",
                     fmt_rate(q["proxy"], q["n"]) if q else "pending",
                     "IP2Location %s / MaxMind %s" % (fmt_rate(mob[p]["ip2l_k"], mob[p]["ip2l_n"]),
                                                      fmt_rate(mob[p]["mm_k"], mob[p]["mm_n"])) if p in mob else "-",
                     None, PRICE[p][0], PRICE[p][1]])
    ws = sheet(wb, "Summary", header, rows,
               widths={1: 14, 2: 22, 3: 14, 4: 16, 5: 11, 6: 22, 7: 9, 8: 13, 9: 22, 10: 22, 11: 12, 12: 26, 13: 22,
                       14: 22, 15: 22, 16: 18, 17: 16, 18: 14, 19: 10, 20: 30},
               pct_cols=(5, 7), num_cols={11: "0.0", 19: "0.00"},
               note="Mobile proxy benchmark, %s. Pooled over three test servers (USA, Germany, Singapore). "
                    "Rates with 95%% Wilson confidence interval. Per-server values in the other tabs." % period)
    for i, p in enumerate(provs):
        ws.cell(4 + i, 18).value = "='Cost'!H%d" % (6 + i)
        ws.cell(4 + i, 18).number_format = "0.000"
    for r in ws.iter_rows(min_row=4, max_row=3 + len(rows)):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r[0].font = Font(bold=True)

    # ---------------- Infra by server
    rows = []
    for p in provs:
        for sv in SERVERS + ["all"]:
            s = P[p] if sv == "all" else PS[(p, sv)]
            if s["n"]:
                rows.append([LABEL[p], SERVER_NAMES.get(sv, "All servers"), s["n"], s["ok"], s["success"], s["lo"], s["hi"],
                             s["median"], s["p90"], s["unique"], s["dup"], s["geo"], s["v6"]])
    sheet(wb, "Infra by server", ["Provider", "Server", "Requests", "Successful", "Success Rate", "CI low", "CI high",
                                  "Latency median, s", "Latency P90, s", "Unique IPs", "Duplicate Rate", "Geo Accuracy",
                                  "IPv6 share"], rows,
          widths={1: 14, 2: 22}, pct_cols=(5, 6, 7, 11, 12, 13), num_cols={8: "0.00", 9: "0.00"},
          note="Infrastructure test: https://www.cloudflare.com/cdn-cgi/trace, new IP on every request, 30 s timeout.")

    # ---------------- Errors
    if len(infra):
        e = infra[infra["ok"] != 1].groupby(["provider", "err"]).size().reset_index(name="count")
        rows = [[LABEL[r.provider], r.err, int(r["count"]), int(r["count"]) / P[r.provider]["n"]]
                for _, r in e.sort_values(["provider", "count"], ascending=[True, False]).iterrows() if r.provider in P]
        sheet(wb, "Infra errors", ["Provider", "Error type", "Count", "Share of all requests"], rows,
              widths={1: 14, 2: 40}, pct_cols=(4,))

    # ---------------- GEO
    if len(infra):
        rows = []
        for p in provs:
            for cc in COUNTRIES:
                g = infra[(infra["provider"] == p) & (infra["country_req"] == cc)]
                if len(g):
                    s = infra_stats(g)
                    rows.append([LABEL[p], cc.upper(), s["n"], s["success"], s["lo"], s["hi"], s["median"], s["p90"],
                                 s["geo"], s["unique"]])
        sheet(wb, "Performance by GEO", ["Provider", "Country", "Requests", "Success Rate", "CI low", "CI high",
                                         "Latency median, s", "Latency P90, s", "Geo Accuracy", "Unique IPs"], rows,
              widths={1: 14, 2: 9}, pct_cols=(4, 5, 6, 9), num_cols={7: "0.00", 8: "0.00"}, freeze="C2")
        for title, col in (("GEO Success pivot", "ok"), ("GEO Accuracy pivot", "geo_ok")):
            rows = []
            for p in provs:
                row = [LABEL[p]]
                for cc in COUNTRIES:
                    g = infra[(infra["provider"] == p) & (infra["country_req"] == cc)]
                    if col == "geo_ok":
                        g = g[g["ok"] == 1]
                    row.append(float(g[col].mean()) if len(g) else None)
                rows.append(row)
            sheet(wb, title, ["Provider"] + [c.upper() for c in COUNTRIES], rows, pct_cols=tuple(range(2, 14)),
                  widths={i: 9 for i in range(2, 14)},
                  note="%s by requested country, pooled over three servers." % ("Success Rate" if col == "ok" else "Geo Accuracy"))

    # ---------------- Targets
    if TG:
        rows = []
        for p in provs:
            for t in ("google", "instagram"):
                for sv in SERVERS + ["all"]:
                    g = targets[(targets["provider"] == p) & (targets["target"] == t)]
                    restricted = int((g["result"] == "provider_restricted").sum())
                    if sv != "all":
                        g = g[g["server"] == sv]
                        restricted = int((g["result"] == "provider_restricted").sum())
                    s = target_stats(g)
                    if not s["n"]:
                        continue
                    lo, hi = wilson(s["k"], s["n"])
                    rows.append([LABEL[p], t, SERVER_NAMES.get(sv, "All servers"), s["n"], s["k"], s["k"] / s["n"], lo, hi,
                                 s["cap"] / s["n"], s["blk"] / s["n"], s["oth"] / s["n"], restricted, s["med"], s["p90"], s["kb"]])
        sheet(wb, "Targets", ["Provider", "Target", "Server", "Requests", "Success", "Target Success Rate", "CI low",
                              "CI high", "CAPTCHA", "Blocked / login wall / rate limited", "Other errors",
                              "Refused by the proxy (excluded)", "Latency median, s", "Latency P90, s", "Avg page, KB (decoded)"],
              rows, widths={1: 14, 2: 11, 3: 22}, pct_cols=(6, 7, 8, 9, 10, 11), num_cols={13: "0.00", 14: "0.00", 15: "0.0"},
              note="Google search (12 countries) and the Instagram profile page of @instagram (US, DE, GB) as a phone browser "
                   "(curl_cffi, Chrome 131 Android TLS fingerprint), no JavaScript. Google 'turn on JavaScript' and consent pages count as let in.")

    # ---------------- Sticky
    rows = []
    for p in provs:
        for sv in SERVERS + ["all"]:
            s = ST[p] if sv == "all" else STS[(p, sv)]
            if s["n"]:
                rows.append([LABEL[p], SERVER_NAMES.get(sv, "All servers"), s["n"], s["med"], s["r10"], s["rf"], s["lo"],
                             s["hi"], s["err"]])
    sheet(wb, "Sticky", ["Provider", "Server", "Sessions", "Median IP lifetime, min", "Lived >= 10 min", "Lived 30 min",
                         "CI low (30 min)", "CI high (30 min)", "Failed checks"], rows,
          widths={1: 14, 2: 22}, pct_cols=(5, 6, 7, 8, 9), num_cols={4: "0.0"},
          note="8 sessions per run asking for a 30-minute sticky IP, one check every 30 s. The lifetime ends at the first other IP "
               "or at 3 failed checks in a row; an IP change at or after 30 minutes counts as kept.")
    if stk:
        sruns = sorted(set(s["run"] for s in stk))
        rows = []
        for p in provs:
            row = [LABEL[p]]
            for run in sruns:
                x = [s for s in stk if s["provider"] == p and s["run"] == run]
                row.append(sum(s["full"] for s in x) / len(x) if x else None)
            rows.append(row)
        sheet(wb, "Trend Sticky by run", ["Provider"] + sruns, rows, widths={j: 14 for j in range(2, len(sruns) + 2)},
              pct_cols=tuple(range(2, len(sruns) + 2)),
              note="Share of sessions that kept the IP 30 minutes, per run (24 sessions per provider per run). Small samples: a trend, not a ranking.")

    # ---------------- Fingerprint
    if FP:
        rows = []
        for p in provs:
            f = FP[p]
            if not f["n"]:
                continue
            n = f["n"]
            rows.append([LABEL[p], f["req"], n, f["linux"] / n, f["windows"] / n, f["mac"] / n, f["other"] / n, f["unknown"] / n,
                         f["eth"] / f["link_n"] if f["link_n"] else None, f["tunnel"] / f["link_n"] if f["link_n"] else None])
        sheet(wb, "Fingerprint", ["Provider", "Requests", "Answers", "Linux (incl. Android)", "Windows", "Mac OS X / iOS",
                                  "Other", "Unknown", "Ethernet MTU 1500", "Tunnel-sized MTU (typical mobile)"], rows,
              widths={1: 14}, pct_cols=(4, 5, 6, 7, 8, 9, 10),
              note="p0f 3 on our own HTTPS endpoint (de server); the exit device opens the TCP connection. Checked with a real iPhone "
                   "on 5G (Mac OS X, tunnel-sized MTU) and on home Wi-Fi (Mac OS X, DSL). Android check pending.")

    # ---------------- IP quality
    if ipq:
        rows = []
        for p in provs:
            q = ipq.get(p)
            if q:
                lo, hi = wilson(q["high"], q["n"])
                rows.append([LABEL[p], q["n"], q["avg"], q["med"], q["high"] / q["n"], lo, hi, q["proxy"] / q["n"],
                             q["dc"] / q["n"], q["vpn"] / q["n"], q["host"] / q["all"], q["geo"] / q["geo_n"] if q["geo_n"] else None])
        sheet(wb, "IP quality", ["Provider", "IPs checked", "Fraud Score avg (0-100)", "Fraud Score median", "High risk share",
                                 "CI low", "CI high", "Proxy flags", "Datacenter", "VPN", "Hosting ASN (name heuristic)",
                                 "Geo accuracy (IPinfo)"], rows, pct_cols=(5, 6, 7, 8, 9, 10, 11, 12),
              num_cols={3: "0.0", 4: "0.0"}, widths={1: 14},
              note="Scamalytics API v3 and IPinfo Lite on a sample of unique exit IPs per provider, spread evenly over the 12 countries.")

    # ---------------- Mobile IP Share
    if mob:
        rows = []
        for p in provs:
            m = mob.get(p)
            if not m:
                continue
            l2, h2 = wilson(m["ip2l_k"], m["ip2l_n"])
            lm, hm = wilson(m["mm_k"], m["mm_n"])
            rows.append([LABEL[p], m["ip2l_n"], m["ip2l_k"] / m["ip2l_n"], l2, h2, m["mm_n"], m["mm_k"] / m["mm_n"], lm, hm,
                         m["both_k"] / m["both_n"], m["mobonly"] / m["both_n"], m["ip2l_not_mm"] / m["both_n"]])
        sheet(wb, "Mobile IP Share", ["Provider", "IPs (IP2Location)", "Mobile by IP2Location", "CI low", "CI high",
                                      "IPs (MaxMind)", "Cellular by MaxMind", "CI low", "CI high", "Mobile in both",
                                      "IP2Location 'MOB' only", "IP2Location mobile, MaxMind not"], rows,
              widths={1: 14}, pct_cols=(3, 4, 5, 7, 8, 9, 10, 11, 12),
              note="Unique exit IPs, up to 10,000 per provider spread over the 12 countries. IP2Location.io: usage type contains "
                   "MOB (also 'ISP/MOB', which means a carrier with both fixed and mobile networks). MaxMind GeoIP2 Insights: "
                   "user type cellular, on the first 1,000 IPs per provider; the last three columns use these 1,000. "
                   "The two databases disagree for some providers; neither is ground truth.")

    # ---------------- Shared IPs
    if shared_any:
        rows = [[LABEL[p], shared_any[p][1], shared_any[p][0], shared_any[p][0] / shared_any[p][1]] for p in provs]
        ws = sheet(wb, "Shared IPs", ["Provider", "Unique exit IPs", "Also given out by another provider", "Share"], rows,
                   widths={1: 14, 2: 14, 3: 18}, pct_cols=(4,),
                   note="Successful infrastructure requests of the whole test. The same public IP at two providers can be "
                        "carrier-grade NAT (many phones behind one address) or the same device in both pools.")
        r0 = len(rows) + 5
        for j, h in enumerate(["Provider A", "Provider B", "Shared IPs", "Share of A's IPs", "Share of B's IPs",
                               "Seen at both in the same hour"], 1):
            c = ws.cell(r0, j, h)
            c.fill, c.font = HEAD_FILL, HEAD_FONT
        for i, (pa, pb, n, fa, fb, same) in enumerate(shared_pairs[:20], 1):
            for j, v in enumerate([LABEL[pa], LABEL[pb], n, fa, fb, same], 1):
                c = ws.cell(r0 + i, j, v)
                if j >= 4:
                    c.number_format = "0.0%"

    # ---------------- Cost (inputs: provider panel counters)
    ws = wb.create_sheet("Cost")
    ws["A1"] = ("Cost per 1,000 successful requests = price per GB x traffic by the provider's own counter / successful requests "
                "of all tests x 1,000. Fill column D with the dashboard counters read right after the test.")
    ws["A1"].font = Font(italic=True, color="555555")
    times = sorted(c["read_at_utc"] for c in counters.values() if c.get("read_at_utc"))
    ws["A2"] = "Counters read at (UTC): " + ("%s ... %s" % (times[0], times[-1]) if times else "")
    ws["A2"].font = Font(bold=True)
    hdr = ["Provider", "Price, $/GB", "Minimum purchase", "Traffic used, MB (dashboard counter)", "Successful infra",
           "Successful targets", "Successful sticky checks + fingerprint", "Cost per 1,000 successful, $",
           "Counter source"]
    cost = {}
    for j, h in enumerate(hdr, 1):
        c = ws.cell(5, j, h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[5].height = 48
    for i, p in enumerate(provs):
        r = 6 + i
        tk = int(targets[(targets["provider"] == p)]["ok"].sum()) if len(targets) else 0
        sk = ST[p]["okreq"] + (FP[p]["n"] if FP else 0)
        c = counters.get(p) or {}
        mb = float(c["used_mb"]) if c.get("used_mb") else None
        for j, v in enumerate([LABEL[p], PRICE[p][0], PRICE[p][1], mb, P[p]["ok"], tk, sk], 1):
            ws.cell(r, j, v)
        # 1 GB = 1,000 MB, as in the residential benchmark (some dashboards count 1,024: a few percent)
        ws.cell(r, 8, '=IF(ISNUMBER(D%d),B%d*D%d/1000/(E%d+F%d+G%d)*1000,"")' % (r, r, r, r, r, r))
        ws.cell(r, 8).number_format = "0.000"
        ws.cell(r, 9, c.get("source", ""))
        if mb is not None:
            cost[p] = PRICE[p][0] * mb / 1000 / (P[p]["ok"] + tk + sk) * 1000
    for j, w in enumerate([14, 11, 36, 18, 14, 14, 18, 16, 60], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B6"

    # ---------------- Runs
    if len(infra):
        rows = []
        for (p, sv, run), x in infra.groupby(["provider", "server", "run_id"]):
            if p in P:
                s = infra_stats(x)
                rows.append([LABEL[p], SERVER_NAMES.get(sv, sv), run, x["ts"].min(), s["n"], s["success"], s["median"]])
        sheet(wb, "Infra runs", ["Provider", "Server", "Run", "Started (UTC)", "Requests", "Success Rate", "Latency median, s"],
              rows, widths={1: 14, 2: 22, 3: 16, 4: 20}, pct_cols=(6,), num_cols={7: "0.00"},
              note="One row per provider, server and scheduled run: shows whether results are stable during the test.")

    # ---------------- Method
    ws = wb.create_sheet("Method and limits")
    text = [
        "Method and limitations (read before citing the numbers)",
        "",
        "Period: %s. Three Hetzner Cloud servers: USA (Ashburn), Germany (Falkenstein), Singapore." % period,
        "11 rotating pay-per-GB mobile proxy providers, 12 countries (US, DE, GB, FR, ES, BR, IN, PL, UA, TR, VN, ID).",
        "Not tested: Rayobyte (sales only, $250 minimum), Infatica (identity verification), Proxy-Cheap (phone and postal address "
        "required), Geonode and Webshare (no mobile proxies), Proxidize and Proxy-Seller (dedicated ports billed by time).",
        "Floppydata is the publisher and one of the tested providers; its traffic came from a package bought for the residential test. "
        "All scripts and raw data are published.",
        "Infrastructure: Cloudflare trace, new IP on every request, 30 s timeout, 5 parallel requests; failures count as results.",
        "Targets: Google search and the Instagram profile page as a phone browser without JavaScript. Success means the site let the IP in.",
        "Sticky: 30-minute sessions, one check every 30 s; IP lifetime ends at another IP or at 3 failed checks in a row.",
        "Fingerprint: p0f on our own endpoint; p0f reports Android as Linux and iOS as Mac OS X; carrier middleboxes can change TCP.",
        "IP quality: Scamalytics and IPinfo Lite on a sample of exit IPs; Mobile IP Share from IP2Location.io and MaxMind "
        "GeoIP2 Insights, shown side by side because the two databases disagree for some providers.",
        "Prices: regular price of the smallest package without promo codes or VAT, checked 23-26 Oct 2026; re-check before publication.",
        "Differences smaller than the width of the confidence intervals are not significant: overlapping intervals are a tie.",
    ]
    for i, t in enumerate(text, 1):
        ws.cell(i, 1, t)
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 160
    for r in ws.iter_rows():
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    if a.json:
        site = {"period": period, "generated_from": "analyze.py", "providers": {}}
        for p in provs:
            s, st = P[p], ST[p]
            g, ig = TG.get((p, "google"), {"k": 0, "n": 0}), TG.get((p, "instagram"), {"k": 0, "n": 0})
            f, q = FP.get(p), ipq.get(p)
            r1 = lambda v, d=1: None if v is None or (isinstance(v, float) and math.isnan(v)) else round(v, d)
            site["providers"][LABEL[p]] = dict(
                successRate=r1(100 * s["success"], 2), requests=s["n"],
                latencyMedian=r1(s["median"], 2), latencyP90=r1(s["p90"], 2),
                latencyByServer=[r1(PS[(p, sv)]["median"], 2) for sv in SERVERS],
                successByServer=[r1(100 * PS[(p, sv)]["success"], 2) for sv in SERVERS],
                uniqueIps=s["unique"], duplicateRate=r1(100 * s["dup"]), geoAccuracy=r1(100 * s["geo"], 2),
                ipv6Share=r1(100 * s["v6"]),
                weakestCountry=weakest.get(p, " ").split(" ")[0], weakestRate=r1(float(weakest.get(p, "x 0%").split(" ")[1].rstrip("%"))),
                googleRate=r1(100 * g["k"] / g["n"]) if g["n"] else None,
                instagramRate=r1(100 * ig["k"] / ig["n"]) if ig["n"] else None,
                stickyMedian=r1(st["med"]), sticky10=r1(100 * st["r10"]), sticky30=r1(100 * st["rf"]), stickySessions=st["n"],
                fpWindows=r1(100 * f["windows"] / f["n"]) if f and f["n"] else None,
                fpEthernet=r1(100 * f["eth"] / f["link_n"]) if f and f["link_n"] else None,
                fraudAvg=r1(q["avg"]) if q else None, fraudHighRisk=r1(100 * q["high"] / q["n"]) if q else None,
                proxyFlags=r1(100 * q["proxy"] / q["n"]) if q else None, ipSample=q["n"] if q else 0,
                mobileShareIp2l=r1(100 * mob[p]["ip2l_k"] / mob[p]["ip2l_n"]) if p in mob else None,
                mobileShareMaxmind=r1(100 * mob[p]["mm_k"] / mob[p]["mm_n"]) if p in mob else None,
                sharedIpShare=r1(100 * shared_any[p][0] / shared_any[p][1]) if p in shared_any else None,
                sharedIpCi=[r1(100 * x, 2) for x in wilson(*shared_any[p])] if p in shared_any else None,
                googleCi=[r1(100 * x, 2) for x in wilson(g["k"], g["n"])] if g["n"] else None,
                googleCaptcha=r1(100 * g["cap"] / g["n"]) if g["n"] else None,
                googleRequests=g["n"],
                costPer1k=r1(cost.get(p), 3), price=PRICE[p][0], minPurchase=PRICE[p][1])
        site["sharedPairs"] = [dict(a=LABEL[pa], b=LABEL[pb], ips=n, shareA=r1(100 * fa), shareB=r1(100 * fb),
                                    sameHour=r1(100 * same)) for pa, pb, n, fa, fb, same in shared_pairs[:6]]
        with open(a.json, "w") as fh:
            json.dump(site, fh, indent=1)
        print("Saved", a.json)

    wb.save(a.out)
    print("Saved", a.out, "| period", period, "| infra rows", len(infra), "| targets rows", len(targets),
          "| sticky sessions", len(stk), "| fingerprint rows", len(fp), "| ip quality providers", len(ipq))


if __name__ == "__main__":
    main()
